#!/usr/bin/env python3
"""Clean the Halo First Thursday guest list.

Reads the Google Form responses and checks each name against the resident
lists (Rezmin rooms tab, Paxton 20V, Paxton ROM). Writes an .xlsx with:
  Guest List      - names that passed every check
  Excluded        - names removed, with the reason
  Check Manually  - close-but-not-exact matches (kept on the list)
  New Sign-ups    - responses not seen on the previous run

Usage:
  python3 guestlist.py --responses responses.csv \
      --db "Rezmin=rezmin.xlsx" --db "Paxton 20V=20v.csv" --db "Paxton ROM=rom.csv" \
      [--exclude-if not-found|found] [--state state.json] [--out out_dir]
"""
import argparse
import csv
import json
import re
import unicodedata
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

FUZZY_THRESHOLD = 0.88


def read_table(path, sheet_hint=None):
    """Return a list of row dicts from a CSV or XLSX file."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = wb.worksheets
        if sheet_hint:
            hinted = [ws for ws in sheets if sheet_hint in ws.title.lower()]
            sheets = hinted or sheets
        rows = []
        for ws in sheets:
            data = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
            rows += rows_from_grid(data)
        return rows
    with open(path, newline="", encoding="utf-8-sig") as f:
        return rows_from_grid(list(csv.reader(f)))


def rows_from_grid(grid):
    """Use the first row that looks like a header (has a name-like column)."""
    for i, row in enumerate(grid[:20]):
        if any(re.search(r"name|surname|guest", (c or "").lower()) for c in row):
            header = [(c or "").strip() for c in row]
            return [dict(zip(header, r)) for r in grid[i + 1:] if any((c or "").strip() for c in r)]
    return []


FILLER = {"yes", "no", "please"}


def normalise(name):
    name = re.sub(r"['’‘`]", "", name or "")  # O'Flaherty -> oflaherty, on both sides
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return [t for t in re.sub(r"[^a-z ]", " ", name.lower()).split() if t not in FILLER]


def split_people(raw):
    """'YES - A B, C D and E F' -> ['A B', 'C D', 'E F']."""
    parts = re.split(r",|;|&|/|\band\b|\n", raw or "", flags=re.I)
    people = []
    for p in parts:
        tokens = normalise(p)
        if tokens:
            people.append(" ".join(t.capitalize() for t in tokens))
    return people


def find_col(headers, *patterns, exclude=()):
    for p in patterns:
        for h in headers:
            hl = h.lower()
            if re.search(p, hl) and not any(re.search(x, hl) for x in exclude):
                return h
    return None


def db_names(rows):
    """Pull full names out of a resident export, whatever its column layout."""
    if not rows:
        return []
    headers = list(rows[0].keys())
    first = find_col(headers, r"^first ?name", r"^name$", r"first")
    last = find_col(headers, r"surname", r"last ?name")
    full = find_col(headers, r"full ?name", r"student ?name", r"resident", r"^name$")
    names = []
    for r in rows:
        if first and last and first != last:
            n = f"{r.get(first, '')} {r.get(last, '')}"
        else:
            n = r.get(full or first or last, "")
        if normalise(n):
            names.append(n.strip())
    return names


def match(form_tokens, db_entries):
    """Return ('exact'|'fuzzy'|None, matched name)."""
    joined = " ".join(sorted(form_tokens))
    best, best_name = 0.0, None
    for name, tokens in db_entries:
        # First and last word of the submitted name both appear: handles middle names and order.
        if form_tokens[0] in tokens and form_tokens[-1] in tokens:
            return "exact", name
        # Same surname, first name starts with the same letter (Nick / Nicholas, Rob / Robert).
        if form_tokens[-1] in tokens and any(t[0] == form_tokens[0][0] and t != form_tokens[-1] for t in tokens):
            best, best_name = 1.0, name
            continue
        score = SequenceMatcher(None, joined, " ".join(sorted(tokens))).ratio()
        if score > best:
            best, best_name = score, name
    return ("fuzzy", best_name) if best >= FUZZY_THRESHOLD else (None, None)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--responses", required=True)
    ap.add_argument("--db", action="append", default=[], help='"Label=path" (repeatable)')
    ap.add_argument("--exclude-if", choices=["not-found", "found"], default="not-found")
    ap.add_argument("--state", default="state.json")
    ap.add_argument("--out", default=".")
    args = ap.parse_args()

    dbs = {}
    for spec in args.db:
        label, _, path = spec.partition("=")
        names = db_names(read_table(path, sheet_hint="room"))
        dbs[label] = [(n, set(normalise(n))) for n in names]
        print(f"{label}: {len(names)} names loaded")

    responses = read_table(args.responses)
    headers = list(responses[0].keys()) if responses else []
    ts_col = find_col(headers, r"timestamp")
    name_col = find_col(headers, r"guest ?list", r"name")
    pres_col = find_col(headers, r"attending|pres")

    state_path = Path(args.state)
    seen = set(json.loads(state_path.read_text())) if state_path.exists() else set()

    kept, excluded, check, new = [], [], [], []
    used = set()
    for r in responses:
        raw = (r.get(name_col) or "").strip()
        ts = r.get(ts_col, "")
        pres = r.get(pres_col, "")
        key = f"{ts}|{raw}"
        if key not in seen:
            new.append([raw, ts, pres])
            seen.add(key)
        # A bare "YES" with no name yields no people: not a guest list request.
        for name in split_people(raw):
            reason, found, fuzzy = check_person(name, dbs, args.exclude_if, used)
            if reason:
                excluded.append([name, raw, ts, reason])
                continue
            kept.append([name, ts, pres, found or "Close match only"])
            if fuzzy:
                check.append([name, ts, fuzzy])

    write_xlsx(Path(args.out), kept, excluded, check, new)
    state_path.write_text(json.dumps(sorted(seen)))
    print(f"Responses: {len(responses)} | New: {len(new)} | Kept: {len(kept)} | "
          f"Excluded: {len(excluded)} | Check manually: {len(check)}")


def check_person(name, dbs, exclude_if, used):
    """Return (exclusion reason or None, found-in labels, fuzzy-match note)."""
    tokens = normalise(name)
    if len(tokens) < 2:
        return "Only one name given, cannot verify", "", ""
    key = " ".join(sorted(tokens))
    if key in used:
        return "Duplicate sign-up", "", ""
    used.add(key)

    found, fuzzy = [], []
    for label, entries in dbs.items():
        kind, who = match(tokens, entries)
        if kind == "exact":
            found.append(label)
        elif kind == "fuzzy":
            fuzzy.append(f"{label}: {who}")

    if exclude_if == "not-found" and not found and not fuzzy:
        return "Not found in " + ", ".join(dbs), "", ""
    if exclude_if == "found" and found:
        return "Already in " + ", ".join(found), "", ""
    return None, ", ".join(found), "; ".join(fuzzy)


def write_xlsx(out, kept, excluded, check, new):
    out.mkdir(parents=True, exist_ok=True)
    xlsx = out / f"Halo_Guest_List_{datetime.now().strftime('%Y-%m-%d_%H%M')}.xlsx"
    wb = Workbook()
    sheets = [
        ("Guest List", ["Name", "Signed up", "Attending pres", "Found in"], sorted(kept, key=lambda x: x[0].lower())),
        ("Excluded", ["Name", "Submitted as", "Signed up", "Reason"], excluded),
        ("Check Manually", ["Name", "Signed up", "Closest match"], check),
        ("New Sign-ups", ["Submitted as", "Signed up", "Attending pres"], new),
    ]
    for i, (title, head, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(50, max(12, *(len(str(c.value or "")) + 2 for c in col)))
    wb.save(xlsx)
    print(f"Saved: {xlsx}")


if __name__ == "__main__":
    main()
